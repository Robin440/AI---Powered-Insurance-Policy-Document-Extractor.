from pathlib import Path
from pypdf import PdfReader
from docx import Document


def read_document(path: str) -> list[dict]:
    """Return page/section-aware text chunks as [{'page_number': int|None, 'text': str}]."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        pages = []
        for page_no, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append({"page_number": page_no, "text": text})
        return pages

    if suffix == ".docx":
        doc = Document(str(file_path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
        return [{"page_number": None, "text": text.strip()}] if text.strip() else []

    if suffix == ".txt":
        text = file_path.read_text(encoding="utf-8", errors="replace").strip()
        return [{"page_number": None, "text": text}] if text else []

    raise ValueError("Unsupported file type. Use PDF, DOCX, or TXT.")


def chunk_pages(pages: list[dict], max_chars: int = 12000) -> list[list[dict]]:
    """Group page text without splitting page metadata; split oversized pages safely."""
    chunks, current, current_size = [], [], 0
    for page in pages:
        text = page["text"]
        page_no = page.get("page_number")
        pieces = [text[i:i + max_chars] for i in range(0, len(text), max_chars)] or [""]
        for piece in pieces:
            item = {"page_number": page_no, "text": piece}
            if current and current_size + len(piece) > max_chars:
                chunks.append(current)
                current, current_size = [], 0
            current.append(item)
            current_size += len(piece)
    if current:
        chunks.append(current)
    return chunks
